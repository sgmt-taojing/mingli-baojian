#!/usr/bin/env python3
"""专利材料 Word 生成 — ①代码文档（前后端各1750行）②界面截图文档（28张带功能说明）"""
import os, json, datetime
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

BASE = os.path.join(os.path.dirname(__file__), '..')
DESKTOP = os.path.join(os.path.expanduser('~'), 'Desktop', '命理宝鉴专利材料')
os.makedirs(DESKTOP, exist_ok=True)

GOLD = RGBColor(0xC9, 0xA2, 0x27)
BROWN = RGBColor(0x6B, 0x42, 0x26)
GRAY = RGBColor(0x88, 0x88, 0x88)
CODE_BG = 'F7F3E8'

def set_font(run, size=10.5, bold=False, color=None, mono=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    if color: run.font.color.rgb = color
    if mono:
        run.font.name = 'Courier New'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Courier New')
    else:
        run.font.name = 'PingFang SC'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), 'PingFang SC')

def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    run = p.add_run(text)
    if level == 0:
        set_font(run, 22, True, GOLD)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif level == 1:
        set_font(run, 16, True, GOLD)
        p.paragraph_format.space_before = Pt(18)
    else:
        set_font(run, 13, True, BROWN)
    p.paragraph_format.space_after = Pt(8)
    return p

def add_para(doc, text, size=10.5, bold=False, color=None, align=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_font(run, size, bold, color)
    if align: p.alignment = align
    return p

def add_code_line(doc, lineno, code, shade=True):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    # 行号
    r1 = p.add_run(f'{lineno:>4}  ')
    set_font(r1, 8, False, GRAY, mono=True)
    # 代码
    r2 = p.add_run(code if code.strip() else ' ')
    set_font(r2, 8, False, None, mono=True)
    if shade:
        pPr = p._p.get_or_add_pPr()
        shd = pPr.makeelement(qn('w:shd'), {qn('w:val'): 'clear', qn('w:color'): 'auto', qn('w:fill'): CODE_BG})
        pPr.append(shd)
    return p

def build_code_doc():
    doc = Document()
    for section in doc.sections:
        section.page_width = Cm(21); section.page_height = Cm(29.7)
        section.left_margin = Cm(2); section.right_margin = Cm(2)
        section.top_margin = Cm(2); section.bottom_margin = Cm(2)

    with open(os.path.join(BASE, '.openclaw', 'tmp', 'patent-meta.json'), encoding='utf-8') as f:
        meta = json.load(f)

    # ===== 封面 =====
    add_para(doc, ' ', 12); add_para(doc, ' ', 12)
    add_heading(doc, '易道智鉴（命理宝鉴）智能诊疗系统', 0)
    add_para(doc, '核心逻辑源代码（申请专利用）', 16, True, BROWN, WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, ' ', 12)
    add_para(doc, f'文档生成日期：{datetime.date.today().strftime("%Y年%m月%d日")}', 11, False, GRAY, WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_page_break()

    # ===== 说明页 =====
    add_heading(doc, '一、文档说明', 1)
    add_para(doc, '1. 本代码为平台核心业务逻辑代码，已去除空行与普通注释，保留方法说明注释（JSDoc），便于审查理解。')
    add_para(doc, f'2. 代码总量：后端 {meta["backend_total"]} 行 + 前端 {meta["frontend_total"]} 行 = {meta["backend_total"] + meta["frontend_total"]} 行。')
    add_para(doc, '3. 后端代码来源于 Node.js API 网关与业务引擎（端口 8920），前端代码来源于浏览器端业务逻辑。')
    add_para(doc, '4. 行号为提取后顺序编号，代码行保持原始缩进，未做任何逻辑改动。')

    add_heading(doc, '二、后端代码来源文件', 2)
    for f, n in meta['backend_sources']:
        add_para(doc, f'  · {f} — 提取 {n} 行', 10, False, BROWN)
    add_heading(doc, '三、前端代码来源文件', 2)
    for f, n in meta['frontend_sources']:
        add_para(doc, f'  · {f} — 提取 {n} 行', 10, False, BROWN)

    # ===== 后端代码 =====
    doc.add_page_break()
    add_heading(doc, f'第二部分  后端核心逻辑代码（{meta["backend_total"]} 行）', 1)
    with open(os.path.join(BASE, '.openclaw', 'tmp', 'patent-backend.txt'), encoding='utf-8') as f:
        back_lines = [l.rstrip('\n') for l in f.readlines()]
    for i, code in enumerate(back_lines, 1):
        add_code_line(doc, i, code)

    # ===== 前端代码 =====
    doc.add_page_break()
    add_heading(doc, f'第三部分  前端核心逻辑代码（{meta["frontend_total"]} 行）', 1)
    with open(os.path.join(BASE, '.openclaw', 'tmp', 'patent-frontend.txt'), encoding='utf-8') as f:
        front_lines = [l.rstrip('\n') for l in f.readlines()]
    for i, code in enumerate(front_lines, 1):
        add_code_line(doc, i, code)

    path = os.path.join(DESKTOP, '易道智鉴-核心源代码-3500行.docx')
    doc.save(path)
    print(f'代码文档: {path} ({len(back_lines)}+{len(front_lines)} 行)')

def build_screenshot_doc():
    doc = Document()
    for section in doc.sections:
        section.page_width = Cm(29.7); section.page_height = Cm(21)
        section.left_margin = Cm(1.5); section.right_margin = Cm(1.5)
        section.top_margin = Cm(1.5); section.bottom_margin = Cm(1.5)

    # 封面
    add_heading(doc, '易道智鉴（命理宝鉴）智能诊疗系统', 0)
    add_para(doc, '系统界面功能截图说明（申请专利用）', 16, True, BROWN, WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, ' ', 10)
    add_para(doc, '本图册从系统登录开始，依次展示各核心功能界面，共 28 张截图，每张附功能说明。', 11, False, GRAY, WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, f'生成日期：{datetime.date.today().strftime("%Y年%m月%d日")}', 11, False, GRAY, WD_ALIGN_PARAGRAPH.CENTER)

    manifest = os.path.join(DESKTOP, '截图', '_manifest.json')
    with open(manifest, encoding='utf-8') as f:
        shots = json.load(f)

    for shot in shots:
        if not shot['path'] or not os.path.exists(shot['path']):
            continue
        doc.add_page_break()
        # 标题（图号+页面名）
        p = doc.add_paragraph()
        run = p.add_run(f"{shot['name']} — 功能说明")
        set_font(run, 13, True, GOLD)
        # 图片
        doc.add_picture(shot['path'], width=Cm(26))
        # 功能说明
        p2 = doc.add_paragraph()
        run2 = p2.add_run('功能说明：' + shot['desc'])
        set_font(run2, 11, False, BROWN)

    path = os.path.join(DESKTOP, '易道智鉴-界面截图说明-28张.docx')
    doc.save(path)
    print(f'截图文档: {path} ({len([s for s in shots if s["path"] and os.path.exists(s["path"])])} 张)')

if __name__ == '__main__':
    build_code_doc()
    build_screenshot_doc()
